import pandas as pd
import numpy as np
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

df_w = pd.read_excel('weights.xlsx', header=None)
df_c = pd.read_excel('cpi_data.xlsx', header=None)
df_h = pd.read_excel('cpi_2020_2022.xlsx', header=None)

# === 데이터 추출 ===
individual_items = []
for idx, row in df_w.iterrows():
    if idx < 6:
        continue
    code = str(row[4]).strip() if pd.notna(row[4]) else ''
    name = str(row[5]).strip() if pd.notna(row[5]) else ''
    w22 = row[6] if pd.notna(row[6]) else None
    cl = str(row[0]).strip() if pd.notna(row[0]) else ''
    cm = str(row[1]).strip() if pd.notna(row[1]) else ''
    cs = str(row[2]).strip() if pd.notna(row[2]) else ''
    cn = str(row[3]).strip() if pd.notna(row[3]) else ''
    if not code or w22 is None:
        continue
    if not cl and not cm and not cs and not cn:
        individual_items.append({'code': code, 'name': name, 'weight': w22})

# 2020-2022 품목별 월별 데이터
hist_data = {}
for idx, row in df_h.iterrows():
    if idx == 0:
        continue
    item_str = str(row[1]).strip().replace('\u3000', '').strip()
    parts = item_str.split(' ', 1)
    if len(parts) < 2:
        continue
    name = parts[1]
    d = {}
    for c in range(2, 14):
        d[f'2020_{c-1:02d}'] = float(row[c]) if pd.notna(row[c]) else None
    for c in range(14, 26):
        d[f'2022_{c-13:02d}'] = float(row[c]) if pd.notna(row[c]) else None
    hist_data[name] = d

# CPI 2025.11~2026.04
months_cpi = ['2025.11', '2025.12', '2026.01', '2026.02', '2026.03', '2026.04']
cpi_data = {}
published_total = {}
for idx, row in df_c.iterrows():
    if idx == 0:
        continue
    item_str = str(row[1]).strip().replace('\u3000', '').strip()
    parts = item_str.split(' ', 1)
    if len(parts) < 2:
        continue
    code, name = parts[0], parts[1]
    if code == '0':
        for i, mn in enumerate(months_cpi):
            published_total[mn] = float(row[i + 2]) if pd.notna(row[i + 2]) else None
    else:
        d = {}
        for i, mn in enumerate(months_cpi):
            d[mn] = float(row[i + 2]) if pd.notna(row[i + 2]) else None
        cpi_data[name] = d

# 총지수 2022 연평균
total_2022_avg = np.mean([float(df_h.iloc[1, c]) for c in range(14, 26)])

# === 엑셀 생성 ===
wb = Workbook()
ws = wb.active
ws.title = 'CPI 검산'

# 스타일
header_font_w = Font(bold=True, size=10, color='FFFFFF')
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
result_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
match_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
calc_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
link_fill = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

num2 = '0.00'
num4 = '0.0000'
num6 = '0.000000'

# 컬럼 배치
# A:No B:코드 C:품목명 D:가중치
# E~P: 2020.01~12  Q: 2020평균
# R~AC: 2022.01~12  AD: 2022평균
# AE: 접속계수
# AF~AK: 2025.11~2026.04 지수
# AL~AQ: w*I/avg22 기여분

C_NO = 1
C_CODE = 2
C_NAME = 3
C_W = 4
C_20S = 5    # 2020.01 시작
C_20E = 16   # 2020.12 끝
C_20A = 17   # 2020 평균
C_22S = 18   # 2022.01 시작
C_22E = 29   # 2022.12 끝
C_22A = 30   # 2022 평균
C_LNK = 31   # 접속계수
C_IS = 32    # 지수 시작 (2025.11)
C_IE = 37    # 지수 끝 (2026.04)
C_CS = 38    # 기여분 시작
C_CE = 43    # 기여분 끝

# 행 1: 제목
ws.merge_cells('A1:AQ1')
ws['A1'] = '소비자물가지수(CPI) 검산표'
ws['A1'].font = Font(bold=True, size=14)

# 행 2: 참조값
ws['A2'] = '총지수 2022연평균(발표치):'
ws['A2'].font = Font(bold=True, size=10)
ws['B2'] = total_2022_avg
ws['B2'].number_format = num4
ws['B2'].font = Font(bold=True, size=10, color='FF0000')

# 행 3: 공식
ws['A3'] = 'CPI(t) = [ SUM(w_i * I_i(t) / avg22_i) / 1000 ] * 총지수2022연평균'
ws['A3'].font = Font(italic=True, size=9, color='555555')

# 행 5: 헤더
HR = 5
headers = ['No', '코드', '품목명', '가중치\n(2022)']
for m in range(1, 13):
    headers.append(f'2020.{m:02d}')
headers.append('2020\n연평균')
for m in range(1, 13):
    headers.append(f'2022.{m:02d}')
headers.append('2022\n연평균')
headers.append('접속계수\n(22avg/20avg)')
for mn in months_cpi:
    headers.append(f'{mn}\n지수')
for mn in months_cpi:
    headers.append(f'{mn}\nw*I/avg22')

for c, h in enumerate(headers, 1):
    cell = ws.cell(row=HR, column=c, value=h)
    cell.font = header_font_w
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin_border

ws.row_dimensions[HR].height = 40

# 데이터 행
DR = 6  # data start row

for i, item in enumerate(individual_items):
    r = DR + i
    name = item['name']

    ws.cell(row=r, column=C_NO, value=i + 1).border = thin_border
    ws.cell(row=r, column=C_CODE, value=item['code']).border = thin_border
    ws.cell(row=r, column=C_NAME, value=name).border = thin_border

    c = ws.cell(row=r, column=C_W, value=item['weight'])
    c.number_format = '0.0'
    c.border = thin_border

    # 2020 월별
    hd = hist_data.get(name, {})
    for m in range(1, 13):
        val = hd.get(f'2020_{m:02d}')
        c = ws.cell(row=r, column=C_20S + m - 1, value=val)
        c.number_format = num2
        c.border = thin_border

    # 2020 연평균 수식
    le = get_column_letter(C_20S)
    lp = get_column_letter(C_20E)
    c = ws.cell(row=r, column=C_20A)
    c.value = f'=AVERAGE({le}{r}:{lp}{r})'
    c.number_format = num4
    c.border = thin_border
    c.fill = calc_fill

    # 2022 월별
    for m in range(1, 13):
        val = hd.get(f'2022_{m:02d}')
        c = ws.cell(row=r, column=C_22S + m - 1, value=val)
        c.number_format = num2
        c.border = thin_border

    # 2022 연평균 수식
    lr = get_column_letter(C_22S)
    lac = get_column_letter(C_22E)
    c = ws.cell(row=r, column=C_22A)
    c.value = f'=AVERAGE({lr}{r}:{lac}{r})'
    c.number_format = num4
    c.border = thin_border
    c.fill = calc_fill

    # 접속계수 수식
    l22a = get_column_letter(C_22A)
    l20a = get_column_letter(C_20A)
    c = ws.cell(row=r, column=C_LNK)
    c.value = f'={l22a}{r}/{l20a}{r}'
    c.number_format = num6
    c.border = thin_border
    c.fill = link_fill

    # 2025.11~2026.04 지수
    cd = cpi_data.get(name, {})
    for j, mn in enumerate(months_cpi):
        val = cd.get(mn)
        c = ws.cell(row=r, column=C_IS + j, value=val)
        c.number_format = num2
        c.border = thin_border

    # w*I/avg22 기여분 수식
    lw = get_column_letter(C_W)
    la22 = get_column_letter(C_22A)
    for j in range(6):
        li = get_column_letter(C_IS + j)
        c = ws.cell(row=r, column=C_CS + j)
        c.value = f'={lw}{r}*{li}{r}/{la22}{r}'
        c.number_format = num4
        c.border = thin_border
        c.fill = result_fill

LR = DR + len(individual_items) - 1  # last data row

# === 합계 행 ===
SR = LR + 2
ws.cell(row=SR, column=C_NAME, value='Sigma(w*I/avg22)').font = Font(bold=True)

# 가중치 합계
lw = get_column_letter(C_W)
c = ws.cell(row=SR, column=C_W)
c.value = f'=SUM({lw}{DR}:{lw}{LR})'
c.number_format = '0.0'
c.font = Font(bold=True)
c.border = thin_border

# 기여분 합계
for j in range(6):
    lc = get_column_letter(C_CS + j)
    c = ws.cell(row=SR, column=C_CS + j)
    c.value = f'=SUM({lc}{DR}:{lc}{LR})'
    c.number_format = num4
    c.font = Font(bold=True)
    c.border = thin_border
    c.fill = result_fill

# === CPI 계산 행 ===
CR = SR + 1
ws.cell(row=CR, column=C_NAME, value='계산 CPI').font = Font(bold=True, size=11)

for j in range(6):
    lc = get_column_letter(C_CS + j)
    c = ws.cell(row=CR, column=C_CS + j)
    c.value = f'={lc}{SR}/1000*$B$2'
    c.number_format = num2
    c.font = Font(bold=True, size=11, color='0000FF')
    c.border = thin_border
    c.fill = match_fill

# === 통계청 발표치 행 ===
PR = CR + 1
ws.cell(row=PR, column=C_NAME, value='통계청 발표치').font = Font(bold=True, size=11)

for j, mn in enumerate(months_cpi):
    c = ws.cell(row=PR, column=C_CS + j, value=published_total[mn])
    c.number_format = num2
    c.font = Font(bold=True, size=11, color='FF0000')
    c.border = thin_border
    c.fill = match_fill

# === 차이 행 ===
DFR = PR + 1
ws.cell(row=DFR, column=C_NAME, value='차이').font = Font(bold=True)

for j in range(6):
    lc = get_column_letter(C_CS + j)
    c = ws.cell(row=DFR, column=C_CS + j)
    c.value = f'={lc}{CR}-{lc}{PR}'
    c.number_format = '0.0000'
    c.font = Font(bold=True)
    c.border = thin_border

# === 열 너비 ===
ws.column_dimensions['A'].width = 5
ws.column_dimensions['B'].width = 10
ws.column_dimensions['C'].width = 14
ws.column_dimensions['D'].width = 8
for ci in range(C_20S, C_20E + 1):
    ws.column_dimensions[get_column_letter(ci)].width = 9
ws.column_dimensions[get_column_letter(C_20A)].width = 10
for ci in range(C_22S, C_22E + 1):
    ws.column_dimensions[get_column_letter(ci)].width = 9
ws.column_dimensions[get_column_letter(C_22A)].width = 10
ws.column_dimensions[get_column_letter(C_LNK)].width = 13
for ci in range(C_IS, C_IE + 1):
    ws.column_dimensions[get_column_letter(ci)].width = 10
for ci in range(C_CS, C_CE + 1):
    ws.column_dimensions[get_column_letter(ci)].width = 13

# 틀 고정
ws.freeze_panes = 'E6'

# 저장
output_path = 'CPI_검산표.xlsx'
wb.save(output_path)
print(f'저장 완료: {output_path}')
print(f'품목 수: {len(individual_items)}')
print(f'데이터 행: {DR}~{LR}')
print(f'합계 행: {SR}, CPI 행: {CR}, 발표치 행: {PR}, 차이 행: {DFR}')
